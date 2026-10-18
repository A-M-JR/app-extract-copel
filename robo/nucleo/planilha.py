"""
Leitura da planilha de entrada da Rotina 1.

Colunas (uma linha por UC):
    nome_cliente    -> vira o nome do arquivo PDF baixado (obrigatório)
    uc              -> unidade consumidora a pesquisar no portal (obrigatório)
    mes_referencia  -> mês da fatura a baixar, ex.: 07/2026 (obrigatório)
    ativo           -> S/N. Deixe "N" pra pular a linha sem apagar (opcional; padrão S)

O login/senha NÃO ficam aqui — são únicos e ficam em credenciais.py.
"""
import re
import datetime as _dt
from pathlib import Path

from openpyxl import load_workbook

import config

OBRIGATORIAS = ["nome_cliente", "uc", "mes_referencia"]
COL_UC_ANEEL = "uc_aneel"
COL_VENCIMENTO = "vencimento"
COL_VALOR = "valor"
COL_STATUS = "status"

# nome dos meses em pt-br (abreviado, completo)
_MESES = {
    1: ("jan", "janeiro"), 2: ("fev", "fevereiro"), 3: ("mar", "março"),
    4: ("abr", "abril"), 5: ("mai", "maio"), 6: ("jun", "junho"),
    7: ("jul", "julho"), 8: ("ago", "agosto"), 9: ("set", "setembro"),
    10: ("out", "outubro"), 11: ("nov", "novembro"), 12: ("dez", "dezembro"),
}
_MES_POR_NOME = {}
for _n, (_ab, _full) in _MESES.items():
    _MES_POR_NOME[_ab] = _n
    _MES_POR_NOME[_full] = _n
_MES_POR_NOME["marco"] = 3   # 'março' sem cedilha


def normalizar_mes(valor):
    """Extrai (mes, ano) de datetime, 'jun/26', '06/2026', 'junho de 2026' etc.
    Devolve (None, None) se não reconhecer."""
    if valor is None:
        return (None, None)
    if isinstance(valor, (_dt.datetime, _dt.date)):
        return (valor.month, valor.year)
    s = str(valor).strip().lower()
    # AAAA-MM  (ex.: 2026-06)
    m = re.match(r"^(\d{4})[/\-.](\d{1,2})$", s)
    if m and 1 <= int(m.group(2)) <= 12:
        return (int(m.group(2)), int(m.group(1)))
    # MM/AAAA ou MM/AA  (ex.: 06/2026, 6/26)
    m = re.match(r"^(\d{1,2})[/\-.](\d{2,4})$", s)
    if m and 1 <= int(m.group(1)) <= 12:
        ano = int(m.group(2))
        return (int(m.group(1)), ano + 2000 if ano < 100 else ano)
    # nome/AAAA  (ex.: jun/26, junho de 2026, junho 2026)
    m = re.match(r"^([a-zç]+)[/\-.\s]+(?:de\s+)?(\d{2,4})$", s)
    if m and m.group(1) in _MES_POR_NOME:
        ano = int(m.group(2))
        return (_MES_POR_NOME[m.group(1)], ano + 2000 if ano < 100 else ano)
    return (None, None)


def variantes_mes(mes: int, ano: int) -> list[str]:
    """Formas de escrever o mês pra tentar casar com o que o portal mostrar."""
    ab, full = _MESES[mes]
    aa, mm = f"{ano % 100:02d}", f"{mes:02d}"
    return [f"{mm}/{ano}", f"{mes}/{ano}", f"{mm}/{aa}", f"{ab}/{ano}", f"{ab}/{aa}",
            f"{full}/{ano}", f"{full} de {ano}", f"{full} {ano}", f"{ab} {ano}",
            f"{ab}{aa}", f"{ab} de {ano}"]


def ler_contas(caminho: Path = config.PLANILHA_CONTAS) -> list[dict]:
    """Lê a planilha e devolve uma lista de UCs ativas a processar."""
    caminho = Path(caminho)
    if not caminho.exists():
        raise FileNotFoundError(
            f"Planilha não encontrada: {caminho}\nRode primeiro:  python main.py modelo"
        )

    wb = load_workbook(caminho, data_only=True)
    ws = wb.active

    # mapeia cabeçalho (linha 1) -> índice, sem depender da ordem das colunas
    cabecalho = {}
    for idx, celula in enumerate(ws[1]):
        nome = str(celula.value).strip().lower() if celula.value is not None else ""
        if nome:
            cabecalho[nome] = idx

    faltando = [c for c in OBRIGATORIAS if c not in cabecalho]
    if faltando:
        raise ValueError(
            f"A planilha precisa ter as colunas {faltando}. "
            f"Colunas encontradas: {list(cabecalho)}"
        )

    contas: list[dict] = []
    for num_linha, linha in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        def val(coluna):
            i = cabecalho.get(coluna)
            return linha[i] if i is not None and i < len(linha) else None

        uc = val("uc")
        nome_cliente = val("nome_cliente")
        mes = val("mes_referencia")

        # linha totalmente vazia -> ignora
        if uc is None and nome_cliente is None and mes is None:
            continue

        ativo = str(val("ativo") or "S").strip().upper()
        if ativo in ("N", "NAO", "NÃO", "0", "FALSE"):
            continue

        if not uc or not mes:
            raise ValueError(
                f"Linha {num_linha}: 'uc' e 'mes_referencia' são obrigatórios "
                f"(uc={uc!r}, mes={mes!r})."
            )

        mes_num, mes_ano = normalizar_mes(mes)
        if not mes_num:
            raise ValueError(
                f"Linha {num_linha}: não entendi o mês {mes!r}. "
                f"Use algo como 06/2026 ou jun/2026."
            )

        contas.append({
            "linha":           num_linha,
            "nome_cliente":    (str(nome_cliente).strip() if nome_cliente else str(uc).strip()),
            "uc":              str(uc).strip(),
            "mes_referencia":  f"{mes_num:02d}/{mes_ano}",       # canônico p/ log
            "mes_num":         mes_num,
            "mes_ano":         mes_ano,
            "mes_variantes":   variantes_mes(mes_num, mes_ano),  # p/ casar no portal
        })

    if not contas:
        raise ValueError("Nenhuma UC ativa encontrada na planilha.")
    return contas


def _indices_cabecalho(ws) -> dict[str, int]:
    """Mapeia nome da coluna (minúsculo) -> índice 1-based."""
    cab: dict[str, int] = {}
    for idx, celula in enumerate(ws[1], start=1):
        if celula.value is not None:
            nome = str(celula.value).strip().lower()
            if nome:
                cab[nome] = idx
    return cab


def garantir_colunas_resultado(ws) -> None:
    """Garante colunas de saída na planilha de entrada."""
    cab = _indices_cabecalho(ws)
    prox = ws.max_column + 1
    for col in (COL_UC_ANEEL, COL_VENCIMENTO, COL_VALOR, COL_STATUS):
        if col not in cab:
            ws.cell(row=1, column=prox, value=col)
            prox += 1


def atualizar_linha(caminho: Path, linha: int, uc_aneel: str = "", status: str = "",
                    vencimento: str = "", valor: str = "") -> None:
    """Atualiza colunas de resultado na linha da planilha de entrada."""
    caminho = Path(caminho)
    wb = load_workbook(caminho)
    ws = wb.active
    garantir_colunas_resultado(ws)
    cab = _indices_cabecalho(ws)
    if uc_aneel:
        ws.cell(row=linha, column=cab[COL_UC_ANEEL], value=uc_aneel)
    if vencimento:
        ws.cell(row=linha, column=cab[COL_VENCIMENTO], value=vencimento)
    if valor:
        ws.cell(row=linha, column=cab[COL_VALOR], value=valor)
    if status:
        ws.cell(row=linha, column=cab[COL_STATUS], value=status)
    try:
        wb.save(caminho)
    except PermissionError:
        raise PermissionError(
            f"Não consegui gravar {caminho.name}. Feche o Excel e rode de novo."
        )
