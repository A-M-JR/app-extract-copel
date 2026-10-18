"""
Gera a planilha modelo (contas.xlsx) da Rotina 1, com cabeçalho, exemplo e instruções.

Uso:
    python criar_modelo.py     (ou:  python main.py modelo)
"""
import sys as _sys
from pathlib import Path as _Path
_sys.path.insert(0, str(_Path(__file__).resolve().parents[1]))
import _caminhos  # noqa: E402,F401  (permite rodar este arquivo direto)

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

import config

_COLUNAS_NOVAS = {"nome_cliente", "uc", "mes_referencia"}


def _tem_colunas_novas(caminho) -> bool:
    """True se o arquivo existente já está no formato da Rotina 1."""
    try:
        ws = load_workbook(caminho, read_only=True).active
        cab = {str(c.value).strip().lower() for c in ws[1] if c.value}
        return _COLUNAS_NOVAS.issubset(cab)
    except Exception:
        return False  # não consegui ler (travado/corrompido) -> tratar como "recriar"


def criar_modelo(caminho=None, sobrescrever=False):
    caminho = config.PLANILHA_CONTAS if caminho is None else caminho
    if caminho.exists() and not sobrescrever:
        if _tem_colunas_novas(caminho):
            print(f"[i] Já existe com as colunas certas: {caminho}")
            return caminho
        print("[i] contas.xlsx está no formato antigo — recriando com as colunas da Rotina 1...")

    wb = Workbook()

    # --- Aba de contas ---
    ws = wb.active
    ws.title = "contas"
    colunas = ["nome_cliente", "uc", "mes_referencia", "ativo",
               "uc_aneel", "vencimento", "valor", "status"]
    ws.append(colunas)

    azul = PatternFill("solid", fgColor="1F6FC4")
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = azul
        celula.alignment = Alignment(horizontal="center")

    # linha de exemplo (apague depois de preencher as suas)
    ws.append(["Cliente Exemplo Ltda", "12345678", "07/2026", "S", "", "", "", ""])

    for col, larg in {"A": 34, "B": 16, "C": 16, "D": 8,
                      "E": 16, "F": 14, "G": 12, "H": 18}.items():
        ws.column_dimensions[col].width = larg

    # --- Aba de instruções ---
    doc = wb.create_sheet("leia-me")
    for texto, negrito in [
        ("ROTINA 1 — COMO PREENCHER A PLANILHA 'contas'", True),
        ("", False),
        ("nome_cliente   = vira o NOME do arquivo PDF baixado (obrigatório).", False),
        ("uc             = unidade consumidora a pesquisar no portal (obrigatório).", False),
        ("mes_referencia = mês da fatura a baixar, ex.: 07/2026 (obrigatório).", False),
        ("ativo          = S processa, N pula a linha sem apagar (padrão S).", False),
        ("", False),
        ("O login e a senha (únicos) NÃO ficam aqui — preencha em credenciais.py.", False),
        ("DICA: apague a linha de exemplo antes de rodar de verdade.", False),
    ]:
        doc.append([texto])
        doc.cell(row=doc.max_row, column=1).font = Font(bold=negrito, size=12 if negrito else 11)
    doc.column_dimensions["A"].width = 90

    try:
        wb.save(caminho)
    except PermissionError:
        print(f"[ERRO] Não consegui gravar {caminho.name}.")
        print("       Feche o contas.xlsx no Excel e rode de novo:  robo modelo")
        return None

    print(f"[ok] Planilha modelo criada: {caminho}")
    print("     Preencha a aba 'contas' e o arquivo credenciais.py, depois rode:  robo testar")
    return caminho


if __name__ == "__main__":
    criar_modelo()
