# -*- coding: utf-8 -*-
"""
Relatório do fluxo por E-MAIL: .xlsx + .html com os totais (OK / DUPLICADA /
NÃO ENCONTRADO / ERRO) e a lista de cada fatura (alias, matrícula, mês, cliente,
valor, retorno do CRM, id, mensagem). Salvos na pasta da Extração-copel.
"""
import html as _html
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

import config

# (chave no registro, título da coluna, largura)
COLS = [
    ("alias", "Alias", 8), ("matricula", "Matrícula", 12), ("mes", "Mês", 9),
    ("nome", "Cliente", 30), ("uc_aneel", "UC ANEEL", 14), ("vencimento", "Venc.", 11),
    ("valor", "Valor", 11), ("consumo", "kWh", 8), ("retorno", "Retorno", 16),
    ("id_fatura", "id CRM", 8), ("titulo_gerado", "Título gerado", 13),
    ("mensagem", "Mensagem", 55), ("caixa", "Caixa", 28),
]
_ORDEM = ["OK", "DUPLICADA", "NAO_ENCONTRADO", "ERRO", "SEM_RESPOSTA"]
_COR = {"OK": "#2e7d32", "DUPLICADA": "#1f6feb", "NAO_ENCONTRADO": "#b54708",
        "ERRO": "#b42318", "SEM_RESPOSTA": "#6b7280", "(ensaio)": "#6b7280"}


def totais(registros) -> dict:
    return dict(Counter(r.get("retorno", "") for r in registros))


def gerar(registros: list[dict]) -> tuple[str, str]:
    """Gera xlsx + html e devolve (caminho_xlsx, caminho_html)."""
    ts = datetime.now()
    base = f"resultado_email_{ts:%Y-%m-%d_%H%M%S}"
    pasta = Path(config.PASTA_EXTRACAO)
    pasta.mkdir(parents=True, exist_ok=True)
    xlsx, htmlp = pasta / f"{base}.xlsx", pasta / f"{base}.html"
    tot = totais(registros)

    # ---- XLSX ----
    wb = Workbook()
    ws = wb.active
    ws.title = "Faturas e-mail"
    ws.append([c[1] for c in COLS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E7D32")
    for r in registros:
        ws.append([r.get(c[0], "") for c in COLS])
    for i, c in enumerate(COLS, 1):
        ws.column_dimensions[chr(64 + i)].width = c[2]
    ws.freeze_panes = "A2"
    wb.save(xlsx)

    # ---- HTML ----
    kpis = "".join(
        f'<div style="flex:1;min-width:110px;background:#fff;border:1px solid #e3e8ee;'
        f'border-radius:10px;padding:12px 14px;text-align:center">'
        f'<div style="font-size:26px;font-weight:700;color:{_COR.get(k, "#1f2937")}">{tot.get(k, 0)}</div>'
        f'<div style="font-size:12px;color:#6b7280">{k}</div></div>'
        for k in _ORDEM if tot.get(k))
    cab = "".join(f'<th style="text-align:left;padding:7px;color:#6b7280;font-size:12px">{c[1]}</th>'
                  for c in COLS)
    linhas = ""
    for r in registros:
        tds = ""
        for c in COLS:
            v = _html.escape(str(r.get(c[0], "")))
            cor = _COR.get(r.get("retorno", ""), "#1f2937") if c[0] == "retorno" else "#1f2937"
            peso = "600" if c[0] == "retorno" else "400"
            tds += (f'<td style="padding:6px 7px;border-top:1px solid #eef1f4;font-size:12px;'
                    f'color:{cor};font-weight:{peso}">{v}</td>')
        linhas += f"<tr>{tds}</tr>"
    html = (
        '<!doctype html><html lang="pt-BR"><head><meta charset="utf-8">'
        '<title>Relatório — faturas por e-mail</title></head>'
        '<body style="font-family:Segoe UI,Arial,sans-serif;background:#f1f4f7;margin:0;padding:24px;color:#1f2937">'
        '<h2 style="color:#0b4f6c;margin:0 0 4px">Faturas por e-mail (Energisa)</h2>'
        f'<div style="color:#6b7280;font-size:13px;margin-bottom:16px">{ts:%d/%m/%Y %H:%M} · '
        f'{len(registros)} fatura(s) processada(s)</div>'
        f'<div style="display:flex;gap:10px;flex-wrap:wrap;margin-bottom:18px">{kpis}</div>'
        '<table style="width:100%;border-collapse:collapse;background:#fff;border:1px solid #e3e8ee;'
        f'border-radius:10px;overflow:hidden"><tr style="background:#f7fafc">{cab}</tr>{linhas}</table>'
        '</body></html>')
    htmlp.write_text(html, encoding="utf-8")
    return str(xlsx), str(htmlp)


def resumo_historico(registros: list[dict], rel_html: str = "") -> dict:
    """Entrada pro histórico (entra no resumo diário por e-mail, agrupado por 'concessionaria')."""
    tot = totais(registros)
    ok = tot.get("OK", 0) + tot.get("DUPLICADA", 0)       # conciliadas com o CRM
    return {
        "concessionaria": "ENERGISA (e-mail)", "fonte": "E-mail (Graph)",
        "continuo": True, "total": len(registros), "ok": ok,
        "pendencias": tot.get("NAO_ENCONTRADO", 0), "erros": tot.get("ERRO", 0),
        "duplicadas": tot.get("DUPLICADA", 0), "novas": tot.get("OK", 0),
        "taxa": round(ok / len(registros) * 100, 1) if registros else 0.0,
        "relatorio": rel_html,
    }
